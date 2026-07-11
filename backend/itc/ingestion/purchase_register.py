"""Parse and validate purchase register (.xlsx) uploads.

Per 02_LLD_ITC_Recovery_Agent (section 6):
    class ColumnMapping(BaseModel): ...  # saved per tenant, reused next time
    def parse_register(file_bytes, mapping: ColumnMapping) -> list[RawRow]: ...
    # TEST: one bad row rejects the whole file with a full error report
    # TEST: saved column mapping is reused on the tenant's next upload
    # TEST: duplicate (invoice_no + tax_period + gstin) reported for confirmation

Every tenant's purchase register uses different column headers (Tally,
SAP, manual-Excel, etc. -- see generate_purchase_register.py's
TENANT_PROFILES). ColumnMapping is the tenant-specific translation from
"whatever this tenant calls the column" to the canonical field name this
system operates on internally. It is designed to be persisted (one row
per tenant in the DB) and reused across uploads rather than re-derived
each time -- ingestion/normalise.py or the tenant settings repository is
the natural home for that persistence; this module only consumes it.
"""

from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError

CANONICAL_FIELDS = [
    "vendor_name",
    "vendor_gstin",
    "invoice_number",
    "invoice_date",
    "item_description",
    "taxable_amount",
    "gst_amount",
    "tax_period",
]


class ColumnMapping(BaseModel):
    """Maps this tenant's actual spreadsheet headers to canonical field names.

    Saved per tenant, reused on their next upload (per the LLD test) so a
    tenant only has to map their columns once, not on every file they send.
    """

    tenant_id: str
    # canonical_field -> this tenant's actual header text, e.g.
    # {"vendor_name": "Party Name", "vendor_gstin": "Party GSTIN", ...}
    columns: dict[str, str]

    def header_for(self, canonical_field: str) -> str:
        return self.columns[canonical_field]

    def canonical_for(self, header: str) -> str | None:
        for canonical, actual in self.columns.items():
            if actual == header:
                return canonical
        return None


class RawRow(BaseModel):
    """One validated purchase register row, in canonical field names."""

    vendor_name: str
    vendor_gstin: str
    invoice_number: str
    invoice_date: date
    item_description: str
    taxable_amount: float
    gst_amount: float
    tax_period: str  # MM/YYYY


class PurchaseRegisterValidationError(Exception):
    """Raised when any row fails validation. Carries every error found."""

    def __init__(self, errors: list[str]) -> None:
        self.errors = errors
        super().__init__(
            f"{len(errors)} validation error(s) in purchase register upload"
        )


class ParseRegisterResult(BaseModel):
    """Result of a successful parse: valid rows plus any duplicate groups.

    Duplicates are surfaced for human confirmation rather than silently
    rejected or silently accepted, per the LLD test: "duplicate (invoice_no
    + tax_period + gstin) reported for confirmation." A duplicate might be
    a genuine data-entry mistake, or a legitimate re-issued invoice -- the
    system doesn't get to decide that unilaterally.
    """

    rows: list[RawRow]
    duplicate_groups: list[list[RawRow]]


def _parse_flexible_date(raw: str) -> date:
    """Purchase registers use inconsistent date formats across tenants."""
    raw = str(raw).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognised date format: {raw!r}")


def parse_register(file_bytes: bytes, mapping: ColumnMapping) -> ParseRegisterResult:
    """Validate and parse a tenant's purchase register .xlsx.

    Collects every row-level error before deciding whether to reject the
    file (no partial commit), per the LLD. Duplicate (invoice_number,
    tax_period, vendor_gstin) combinations are valid individually but
    grouped and returned separately for confirmation, not treated as
    errors.
    """
    errors: list[str] = []

    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
    except Exception as exc:
        raise PurchaseRegisterValidationError(
            [f"file could not be read as .xlsx: {exc}"]
        ) from exc

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise PurchaseRegisterValidationError(["file has no header row"]) from exc

    headers = [str(h).strip() if h is not None else "" for h in header_row]

    # Build header-index -> canonical-field lookup using this tenant's
    # saved mapping. Missing expected columns is a hard, whole-file reject
    # (the mapping itself is wrong for this file, not a per-row problem).
    header_to_index = {h: i for i, h in enumerate(headers)}
    field_index: dict[str, int] = {}
    for field in CANONICAL_FIELDS:
        expected_header = mapping.header_for(field)
        if expected_header not in header_to_index:
            errors.append(
                f"expected column {expected_header!r} (for {field!r}) not "
                f"found in file headers: {headers}"
            )
        else:
            field_index[field] = header_to_index[expected_header]

    if errors:
        raise PurchaseRegisterValidationError(errors)

    parsed_rows: list[RawRow] = []
    seen: dict[tuple[str, str, str], list[RawRow]] = {}

    for row_num, row in enumerate(rows_iter, start=2):  # 1-indexed + header row
        if row is None or all(v is None for v in row):
            continue  # skip fully blank rows

        raw_values = {field: row[idx] for field, idx in field_index.items()}
        row_ctx = f"row {row_num}"

        try:
            invoice_date = _parse_flexible_date(raw_values["invoice_date"])
        except (ValueError, TypeError) as exc:
            errors.append(f"{row_ctx}: invalid invoice_date: {exc}")
            continue

        tax_period_raw = str(raw_values["tax_period"]).strip()
        # accept either "MM/YYYY" (as our generator writes) or "MMYYYY"
        if "/" in tax_period_raw:
            tax_period = tax_period_raw
        elif len(tax_period_raw) == 6 and tax_period_raw.isdigit():
            tax_period = f"{tax_period_raw[:2]}/{tax_period_raw[2:]}"
        else:
            errors.append(f"{row_ctx}: invalid tax_period: {tax_period_raw!r}")
            continue

        try:
            raw_row = RawRow(
                vendor_name=str(raw_values["vendor_name"]),
                vendor_gstin=str(raw_values["vendor_gstin"]),
                invoice_number=str(raw_values["invoice_number"]),
                invoice_date=invoice_date,
                item_description=str(raw_values["item_description"]),
                taxable_amount=float(raw_values["taxable_amount"]),
                gst_amount=float(raw_values["gst_amount"]),
                tax_period=tax_period,
            )
        except (ValidationError, ValueError, TypeError) as exc:
            errors.append(f"{row_ctx}: {exc}")
            continue

        parsed_rows.append(raw_row)
        dup_key = (raw_row.invoice_number, raw_row.tax_period, raw_row.vendor_gstin)
        seen.setdefault(dup_key, []).append(raw_row)

    if errors:
        raise PurchaseRegisterValidationError(errors)

    duplicate_groups = [group for group in seen.values() if len(group) > 1]

    return ParseRegisterResult(rows=parsed_rows, duplicate_groups=duplicate_groups)


def mapping_from_tenant_profile(
    tenant_id: str, columns: dict[str, str]
) -> ColumnMapping:
    """Convenience constructor matching generate_purchase_register.py's
    TENANT_PROFILES[tenant_id]["columns"] shape, for wiring the two
    together during development/testing."""
    return ColumnMapping(tenant_id=tenant_id, columns=columns)
